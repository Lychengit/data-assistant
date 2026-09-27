#!/usr/bin/env python3
"""医生月度绩效 -> Excel 工件（技能 doctor_perf_excel_report 的固定脚本）。

本脚本的取数策略是 manifest 里的 ``script.data = none``：**它不联网、不读令牌、不认识任何密钥**。
它只做一件事——把宿主侧准备好的输入工件渲染成 Excel，写到指定路径。

为什么坚持这样切：脚本要是在容器里自己发请求，就得把令牌发进容器；令牌一旦进去，就可能被写进产物
带出系统，平台的限流 / 审计 / 撤销随之全部失效（§18.4.4）。所以"取数"留在平台侧，"渲染"留在沙箱里。

用法（两种输入方式，二选一）::

    # ① 输入已在文件里
    python3 gen_perf_excel.py --input /in/input.json --output /out/report.xlsx

    # ② 数据直接在命令行上（**沙箱里推荐这一种**，见下）
    python3 gen_perf_excel.py --data-base64 '<base64(UTF-8 JSON)>' --output /out/report.xlsx

**为什么推荐 ``--data-base64``**：沙箱命令要经过宿主机的 ``docker exec sh -c <命令>``。
在 Windows 宿主上，命令里的**双引号曾被吃掉**（实测：``open("a.json")`` 变成 ``open(a.json)``，
Python 直接 SyntaxError）——平台已于 2026-09-27 修复：宿主 JVM 会转义内嵌引号。
即便如此，base64 仍是最稳的一条路：只有 ``A-Za-z0-9+/=``，不含引号也不含中文，
**整条命令是纯 ASCII**，不经过任何引号 / 编码环节。

加 ``--emit-base64`` 还会把生成好的 xlsx 以 base64 打进标准输出，
宿主侧可以直接把它交给 ``POST /doctor/export/upload`` 的 ``content_base64``，不必再读文件。

输入工件结构::

    {
      "month": "2026-09",
      "metric_key": "outpatient_visits",
      "metric_name": "门诊人次",
      "unit": "人次",
      "basis": "按挂号口径统计",
      "doctors": [
        {"doctor_id": "1", "doctor_name": "张医生", "dept_code": "心内科", "metric_value": 120}
      ],
      "missing": ["9"]   # 当月没查到绩效的医生编号；**必须是 doctors 的子集**（只为标注，不参与人数统计）
    }

``metric_value`` 允许为 ``null``：名单里有这位医生、但本月没有绩效行，是**正常情况**。
报表里留空并在摘要里写明人数——补 0 会把"没有数据"伪装成"绩效为零"，那是两件事。
"""

from __future__ import annotations

import argparse
import base64
import json
import sys
from datetime import datetime, timezone

HEADERS = ["医生编号", "医生姓名", "科室", "统计月份", "指标", "指标值", "单位"]


def decode_base64_text(encoded: str) -> str:
    """base64 -> UTF-8 文本。解码失败要报清楚是编码问题，别让它伪装成 JSON 语法错。"""
    try:
        raw = base64.b64decode(encoded, validate=True)
    except Exception as error:  # noqa: BLE001 - 只用来把异常换成一句能照着改的话
        raise ValueError(f"--data-base64 不是合法 base64：{error}") from error
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError as error:
        raise ValueError(f"--data-base64 解出来不是 UTF-8 文本：{error}") from error


def load_input(path: str | None, data_base64: str | None) -> dict:
    if data_base64:
        payload = json.loads(decode_base64_text(data_base64))
    elif path:
        with open(path, "r", encoding="utf-8") as handle:
            payload = json.load(handle)
    else:
        raise ValueError("必须给 --input 或 --data-base64 之一")
    if not isinstance(payload, dict):
        raise ValueError("输入工件必须是 JSON 对象")
    payload.setdefault("doctors", [])
    payload.setdefault("missing", [])
    return payload


def build_workbook(payload: dict):
    # openpyxl 在构建镜像时装好（容器 network=none，运行期装不了依赖）。
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    month = payload.get("month") or ""
    metric_key = payload.get("metric_key") or ""
    metric_name = payload.get("metric_name") or metric_key
    unit = payload.get("unit") or ""
    basis = payload.get("basis") or ""
    doctors = payload.get("doctors") or []

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "绩效明细"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    sheet.append(HEADERS)
    for column, _ in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for doctor in doctors:
        value = doctor.get("metric_value")
        sheet.append([
            doctor.get("doctor_id") or "",
            doctor.get("doctor_name") or "",
            doctor.get("dept_code") or "",
            month,
            metric_key,
            value if value is not None else None,
            unit,
        ])

    widths = [12, 14, 12, 12, 18, 12, 8]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    sheet.freeze_panes = "A2"

    notes = workbook.create_sheet("口径说明")
    notes.append(["项目", "值"])
    notes.append(["统计月份", month])
    notes.append(["指标", f"{metric_name}（{metric_key}）"])
    notes.append(["口径", basis])
    notes.append(["单位", unit])
    # 名单人数只数 doctors：missing 是「名单里没查到数据的那几位」的编号，
    # 把它再加一遍会把同一批人算两遍——实测 2026-09-27：6 位医生里有 2 位当月没有绩效行，
    # 报表的「医生数（名单）」却写成了 8。
    no_data_ids = {str(item) for item in (payload.get("missing") or [])}
    no_data_ids |= {str(doctor.get("doctor_id")) for doctor in doctors if doctor.get("metric_value") is None}
    notes.append(["医生数（名单）", len(doctors)])
    notes.append(["无绩效数据的医生数", len(no_data_ids)])
    notes.append(["生成时间(UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    notes.append(["说明", "本表由技能 doctor_perf_excel_report 的固定脚本生成；数据来源见绑定接口的溯源标注"])
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["B"].width = 60

    return workbook


def main() -> int:
    parser = argparse.ArgumentParser(description="医生月度绩效 -> Excel")
    parser.add_argument("--input", help="输入工件 JSON 路径")
    parser.add_argument("--data-base64", dest="data_base64", help="输入工件 JSON 的 base64（UTF-8），沙箱里推荐用这个")
    parser.add_argument("--output", required=True, help="输出 xlsx 路径")
    parser.add_argument("--emit-base64", dest="emit_base64", action="store_true",
                        help="把生成的 xlsx 也以 base64 打进标准输出（供上传接口的 content_base64 直接用）")
    args = parser.parse_args()

    payload = load_input(args.input, args.data_base64)
    workbook = build_workbook(payload)
    workbook.save(args.output)

    # 只往标准输出打一行结构化结果：宿主代理据此核对"确实生成了工件"。
    result = {
        "ok": True,
        "output": args.output,
        "rows": len(payload.get("doctors") or []),
        "month": payload.get("month"),
        "metric_key": payload.get("metric_key"),
    }
    if args.emit_base64:
        with open(args.output, "rb") as handle:
            result["content_base64"] = base64.b64encode(handle.read()).decode("ascii")
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())